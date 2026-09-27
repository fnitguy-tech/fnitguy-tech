"""Build FNIT Guy's 'Asset & Mitigation Checklist' threat modeling workbook.

Usage:  python build_checklist.py Asset-Mitigation-Checklist.xlsx

Content originates from the 2022 workbook for the "Threat Modeling for Normal People"
video, refreshed in 2026 (Verizon 2026 DBIR stats, passkeys/password managers,
Applications + Information checklists, priorities). This script controls structure
and styling so the file can be regenerated.
"""
import math
import sys

from PIL import ImageFont
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

OUT = sys.argv[1]
VIDEO = "https://youtu.be/1WyiLTxV6YI"
SITE = "https://www.fnitguy.tech"
DBIR26 = "https://www.verizon.com/business/resources/reports/2026-dbir-data-breach-investigations-report.pdf"

# ------------------------------------------------------------------ palette
NAVY = "17324D"
NAVY_SOFT = "C9D6E3"
BLUE = "2A78D6"
BLUE_TINT = "EAF2FC"
CARD = "F5F6F8"
INK = "1F2933"
INK2 = "52514E"
MUTED = "898781"
LINE = "E1E0D9"
SERIES = ["2A78D6", "EB6834", "1BAF7A"]   # validated categorical order (blue, orange, aqua)
GOOD = ("E3F4E3", "0B6B0B")
WARN = ("FFF3D6", "8A5A00")
BAD = ("FBE4E4", "A32424")
NEUTRAL = ("EEEEEC", "52514E")
F = "Arial"


def font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


SIDE = Side(style="thin", color=LINE)
ROWLINE = Border(bottom=SIDE)
ACCENT_LEFT = Border(left=Side(style="thick", color=BLUE))
TOP_WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
VCENTER = Alignment(vertical="center")

wb = Workbook()
wb.properties.creator = "Zachary Rogers (FNIT Guy)"
wb.properties.title = "Asset & Mitigation Checklist - Threat Modeling for Normal People"


# ------------------------------------------------------------------ helpers
def col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def setup(ws, widths, tab_color):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = tab_color
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[col(i)].width = w
    ws.last = len(widths)
    ws.px_of = lambda a, b: sum(w * 7 + 5 for w in widths[a - 1:b])   # usable pixels across columns a..b


def banner(ws, title, subtitle, back=True):
    last = ws.last
    for r, h in ((1, 10), (2, 34), (3, 20), (4, 12)):
        ws.row_dimensions[r].height = h
        for c in range(1, last + 1):
            ws.cell(row=r, column=c).fill = fill(NAVY)
    ws.cell(row=2, column=2, value=title).font = font(20, True, "FFFFFF")
    ws.cell(row=2, column=2).alignment = VCENTER
    ws.cell(row=3, column=2, value=subtitle).font = font(10, False, NAVY_SOFT)
    ws.cell(row=3, column=2).alignment = VCENTER
    if back:
        bc = last - 1
        if ws.column_dimensions[col(bc)].width < 12:   # narrow column: widen the button over two columns
            ws.merge_cells(start_row=2, start_column=bc - 1, end_row=2, end_column=bc)
            bc -= 1
            ws.cell(row=2, column=bc + 1).fill = fill(BLUE_TINT)
        c = ws.cell(row=2, column=bc, value="← Start Here")
        c.hyperlink = Hyperlink(ref=c.coordinate, location="'Start Here'!A1", display="Start Here")
        c.font = font(9, True, NAVY)
        c.fill = fill(BLUE_TINT)
        c.alignment = CENTER
    ws.row_dimensions[5].height = 12


# ---- real text measurement (Arial metrics via Pillow) so columns and rows fit their text
_FONTS = {}
_FILES = {(False, False): "arial.ttf", (True, False): "arialbd.ttf", (False, True): "ariali.ttf", (True, True): "arialbi.ttf"}


def _pil(size, bold=False, italic=False):
    key = (size, bold, italic)
    if key not in _FONTS:
        try:
            _FONTS[key] = ImageFont.truetype(r"C:\Windows\Fonts\\" + _FILES[(bold, italic)], round(size * 96 / 72))
        except OSError:
            _FONTS[key] = ImageFont.truetype(_FILES[(bold, italic)], round(size * 96 / 72))
    return _FONTS[key]


def text_px(text, size=10, bold=False, italic=False):
    return _pil(size, bold, italic).getlength(str(text))


def lines_for(text, avail_px, size=10, bold=False, italic=False):
    """Number of wrapped lines `text` needs in a box `avail_px` wide (greedy word wrap, like Excel)."""
    f = _pil(size, bold, italic)
    n = 0
    for para in str(text).split("\n"):
        n += 1
        line = ""
        for word in para.split(" "):
            trial = f"{line} {word}" if line else word
            if not line or f.getlength(trial) <= avail_px:
                line = trial
            else:
                n += 1
                line = word
    return n


def units(px):
    """Pixels -> Excel column-width units (7 px per unit at default font), rounded up to 0.5."""
    return math.ceil((px + 5) / 7 * 2) / 2


def row_h(lines, size=10):
    """Row height in points for `lines` lines of text at `size` pt."""
    return lines * size * 1.42


def plain(text):
    return text if isinstance(text, str) else "".join(t.text if hasattr(t, "text") else str(t) for t in text)


def block(ws, r, text, c1=2, c2=None, size=10, bold=False, color=INK, italic=False,
          bg=None, border=None, pad=6, indent=0, align=TOP_WRAP):
    c2 = c2 or ws.last - 1
    if border is ACCENT_LEFT or bg:          # callouts/cards: keep text off the edge
        indent = max(indent, 2)
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(row=r, column=c1, value=text)
    cell.font = font(size, bold, color, italic)
    has_bold = bold or (not isinstance(text, str) and any(getattr(t.font, "b", False) for t in text))
    avail = ws.px_of(c1, c2) - indent * 10 - 10          # indent level ~10px, plus cell padding
    nlines = lines_for(plain(text), avail, size, has_bold, italic)
    # wrap only when needed: LibreOffice ignores indent on wrapped cells
    cell.alignment = Alignment(wrap_text=nlines > 1, vertical=align.vertical, horizontal=align.horizontal, indent=indent)
    if bg:
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).fill = fill(bg)
    if border:
        ws.cell(row=r, column=c1).border = border
    ws.row_dimensions[r].height = row_h(nlines, size) + pad
    return cell


def section(ws, r, text, c2=None):
    c2 = c2 or ws.last - 1
    ws.row_dimensions[r].height = 26
    cell = ws.cell(row=r, column=2, value=text)
    cell.font = font(13, True, NAVY)
    cell.alignment = Alignment(vertical="bottom")
    for c in range(2, c2 + 1):
        ws.cell(row=r, column=c).border = Border(bottom=Side(style="medium", color=BLUE))
    ws.row_dimensions[r + 1].height = 6
    return r + 2


def spacer(ws, r, h=10):
    ws.row_dimensions[r].height = h
    return r + 1


def bullets(ws, r, items, c2=None, inline=False):
    """inline=True puts the bullet inside the text (for sheets whose column B is wide)."""
    c2 = c2 or ws.last - 1
    for it in items:
        if inline:
            block(ws, r, rich(("•  ", True, BLUE), (it, False, INK)), c1=2, c2=c2, pad=5, indent=1)
        else:
            m = ws.cell(row=r, column=2, value="•")
            m.font = font(10, True, BLUE)
            m.alignment = Alignment(horizontal="right", vertical="top")
            block(ws, r, it, c1=3, c2=c2, pad=5)
        r += 1
    return r


def rich(*parts, size=10):
    out = CellRichText()
    for text, bold, color in parts:
        if text:
            out.append(TextBlock(InlineFont(rFont=F, sz=size, b=bold, color=color), text))
    return out


def link(cell, url, text=None, size=10, bold=False):
    if url.startswith("#"):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=url[1:], display=text or cell.value)
    else:
        cell.hyperlink = url
    if text:
        cell.value = text
    cell.font = font(size, bold, BLUE, underline="single")


def print_setup(ws, landscape=False, rows_to_repeat=None):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if rows_to_repeat:
        ws.print_title_rows = rows_to_repeat


def status_rule(ws, rng, value, style):
    bg, ink = style
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill(bg),
                                                  font=Font(name=F, color=ink, bold=True)))


def tile(ws, row, c1, c2, value, label, fmt="0", big=24):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 1, end_column=c2)
    v = ws.cell(row=row, column=c1, value=value)
    v.number_format = fmt
    v.font = font(big, True, NAVY)
    v.alignment = Alignment(horizontal="left", vertical="bottom", indent=2)
    lab = ws.cell(row=row + 1, column=c1, value=label)
    lab.font = font(9, True, INK2)
    lab.alignment = Alignment(vertical="top", indent=2,
                              wrap_text=lines_for(label, ws.px_of(c1, c2) - 30, 9, True) > 1)
    for rr in (row, row + 1):
        for cc in range(c1, c2 + 1):
            ws.cell(row=rr, column=cc).fill = fill(CARD)
    gap = Side(style="thick", color="FFFFFF")   # white gutter so adjacent tiles read as separate cards
    for cc in range(c1, c2 + 1):
        ws.cell(row=row, column=cc).border = Border(top=Side(style="medium", color=BLUE), left=gap if cc == c1 else None)
    ws.cell(row=row + 1, column=c1).border = Border(left=gap)
    ws.row_dimensions[row].height = 38
    ws.row_dimensions[row + 1].height = max(ws.row_dimensions[row + 1].height or 0,
                                            row_h(lines_for(label, ws.px_of(c1, c2) - 20, 9, True), 9) + 10)


# ================================================================== trackers
FIRST, LAST = 8, 207          # data rows (row 6 header, row 7 example)
YN, YNM, LMH = ["Yes", "No"], ["Yes", "No", "Maybe"], ["Low", "Moderate", "High"]
CAT = {"High": BAD, "Moderate": WARN, "Low": GOOD}
YES_GOOD = {"Yes": GOOD, "No": BAD}
PRIORITY = {"Fix first": BAD, "Fix next": WARN, "Later": NEUTRAL, "Done": GOOD, "Set category": NEUTRAL}


def tracker(ws, title, subtitle, instructions, cols, example, notes):
    """cols: list of (header, min_width, options, {value: status}, bad_value_or_None)."""
    def fit(header, values=(), italic_values=(), min_units=0, pad=18):
        w = max(text_px(word, 10, True) for word in header.split())
        while lines_for(header, w, 10, True) > 2:         # header may wrap, but to two lines at most
            w += 4
        w = max([w] + [text_px(v, 10, True) for v in values] + [text_px(v, 10, False, True) + 12 for v in italic_values])
        return max(min_units, units(w + pad))

    data_w = []
    for i, (h, mn, options, status, bad) in enumerate(cols):
        ex = [str(example[i])] if i < len(example) and example[i] else []
        data_w.append(fit(h, options or [], ex, mn))
    gaps_w = fit("Gaps", ["7 gaps", "None"])
    prio_w = fit("Priority", list(PRIORITY), ["Fix first"])
    widths = [2.5] + data_w + [gaps_w, prio_w, 2.5]
    setup(ws, widths, BLUE)
    banner(ws, title, subtitle)
    ncol = len(cols)
    gap_c, pri_c = ncol + 2, ncol + 3
    G, P = col(gap_c), col(pri_c)
    block(ws, 5, instructions, c1=2, c2=pri_c, color=INK2, pad=10, align=VCENTER)
    ws.row_dimensions[5].height = max(ws.row_dimensions[5].height + 8, 30)
    headers = [c[0] for c in cols] + ["Gaps", "Priority"]
    notes = dict(notes)
    notes.setdefault("Gaps", "Counted automatically: every red answer in the row is one gap.")
    notes.setdefault("Priority", "Fix first = High category with gaps.  Fix next = Moderate.  Later = Low.  "
                                 "Done = no gaps.  Sort or filter this column to see what to tackle first.")
    for i, h in enumerate(headers, start=2):
        c = ws.cell(row=6, column=i, value=h)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = CENTER if i > 2 else Alignment(vertical="center", indent=2)
        c.border = Border(right=Side(style="thin", color="2E4A66"))
        if h in notes:
            c.comment = Comment(notes[h], "FNIT Guy", width=300, height=150)
    head_lines = max(lines_for(h, ws.px_of(i, i) - 12, 10, True) for i, h in enumerate(headers, start=2))
    ws.row_dimensions[6].height = max(30, row_h(head_lines) + 12)
    for r in range(7, LAST + 1):
        ws.row_dimensions[r].height = 20
        for i in range(2, pri_c + 1):
            c = ws.cell(row=r, column=i)
            c.border = ROWLINE
            c.font = font(10)
            c.alignment = Alignment(vertical="center", horizontal="center" if i > 2 else "left",
                                    indent=0 if i > 2 else 2)
    for i, v in enumerate(example, start=2):
        ws.cell(row=7, column=i, value=v)
    bad_terms = [(col(i), c[4]) for i, c in enumerate(cols, start=2) if c[4]]
    for r in range(7, LAST + 1):
        terms = "+".join(f'({L}{r}="{bad}")' for L, bad in bad_terms)
        g = ws.cell(row=r, column=gap_c, value=f'=IF(B{r}="","",{terms})')
        g.number_format = '[=0]"None";[=1]"1 gap";0" gaps"'
        p = ws.cell(row=r, column=pri_c,
                    value=f'=IF(B{r}="","",IF({G}{r}=0,"Done",IF(C{r}="High","Fix first",'
                          f'IF(C{r}="Moderate","Fix next",IF(C{r}="Low","Later","Set category")))))')
        for c in (g, p):
            c.font = font(10, True, INK)
    for i in range(2, pri_c + 1):
        c = ws.cell(row=7, column=i)
        c.font = font(10, i >= gap_c, MUTED, italic=True)
        c.fill = fill(CARD)
    for i, (h, w, options, status, bad) in enumerate(cols, start=2):
        L = col(i)
        if options:
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Pick from the list",
                                error="Choose one of: " + ", ".join(options))
            ws.add_data_validation(dv)
            dv.add(f"{L}7:{L}{LAST}")
        for value, style in status.items():
            status_rule(ws, f"{L}{FIRST}:{L}{LAST}", value, style)
    rng = f"{G}{FIRST}:{G}{LAST}"
    for cond, style in ((f"{G}{FIRST}>=3", BAD), (f"{G}{FIRST}>=1", WARN), (f"{G}{FIRST}=0", GOOD)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({G}{FIRST}),{cond})"], fill=fill(style[0]),
                                                       font=Font(name=F, color=style[1], bold=True)))
    for value, style in PRIORITY.items():
        status_rule(ws, f"{P}{FIRST}:{P}{LAST}", value, style)
    ws.freeze_panes = "C7"
    ws.auto_filter.ref = f"B6:{P}{LAST}"
    print_setup(ws, landscape=True, rows_to_repeat="6:6")
    ws.print_area = f"A1:{col(pri_c + 1)}60"
    return {"name": ws.title, "gaps": G, "priority": P}


# ------------------------------------------------------------------ sheet order
start = wb.active
start.title = "Start Here"

# ================================================================== Definitions
ws = wb.create_sheet("Definitions")
setup(ws, [2.5, 22, 72, 22, 2.5], NAVY)
banner(ws, "Definitions", "The words you’ll see throughout this workbook")
r = 6
for i, h in enumerate(["Term", "Definition", "Source"], start=2):
    c = ws.cell(row=r, column=i, value=h)
    c.font = font(10, True, "FFFFFF")
    c.fill = fill(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[r].height = 24
r += 1
NIST = "NIST glossary"
terms = [
    ("The basics", None, None),
    ("Threat", "Anyone or anything that could cause harm: a thief, a scammer, malware, even a flood.", ""),
    ("Vulnerability", "A weakness a threat can use: an unpatched app, a reused password, an unlocked door.", ""),
    ("Risk", "Threat + vulnerability. If both exist, you have risk – and controls are how you reduce it.", ""),
    ("Attack surface", "The sum of all the points where an unauthorized user can access a system.", ""),
    ("Attack vector", "The means by which a threat gains access to a point of entry.", ""),
    ("Social engineering", None, None),
    ("Social engineering", "A general term for attackers trying to trick people into revealing sensitive information or "
     "performing certain actions, such as downloading and executing files that appear to be benign but are actually malicious.", ""),
    ("Phishing", "Tricking individuals into disclosing sensitive personal information through deceptive computer-based means.",
     "NIST SP 800-83"),
    ("", "Deceiving individuals into disclosing sensitive personal information through deceptive computer-based means.",
     "CNSSI-4009"),
    ("", "A digital form of social engineering that uses authentic-looking—but bogus—emails to request information "
     "from users or direct them to a fake Web site that requests information.", "NIST SP 800-115"),
    ("Vishing", "See phishing, but with voice (phone calls).", ""),
    ("Smishing", "See phishing, but via SMS (text messages).", ""),
    ("Pretexting", "Building a fake but believable scenario (“this is IT support”, “this is your bank”) to "
     "trick someone into taking an action, frequently by phone.", "Verizon 2026 DBIR"),
    ("Protecting yourself", None, None),
    ("MFA", "Multi-factor authentication: a second proof on top of your password, like an authenticator app code, a security "
     "key or a passkey.", ""),
    ("Passkey", "A replacement for passwords that uses your phone or computer (and its fingerprint, face or PIN) to sign you "
     "in. Passkeys can’t be phished, guessed or reused.", "FIDO Alliance"),
    ("Password manager", "An app that creates, stores and fills a unique, strong password for every account, so you only "
     "have to remember one.", "CISA"),
    ("STIG", "Security Technical Implementation Guide: very specific hardening instructions for a product, published by "
     "DISA for the U.S. Department of Defense.", "DISA"),
    ("Zero-day", "A vulnerability attackers are using before the vendor has released a fix.", ""),
    ("Ransomware", "Malware that locks or steals your files and demands payment. Good, tested backups take away its power.", ""),
    ("End of life (EOL)", "A product the vendor no longer supports with security updates. It will never be patched again.", ""),
]
band = False
for term, definition, source in terms:
    if definition is None:
        r = spacer(ws, r, 8)
        g = ws.cell(row=r, column=2, value=term.upper())
        g.font = font(9, True, BLUE)
        g.alignment = Alignment(vertical="bottom", indent=1)
        ws.row_dimensions[r].height = 20
        for cc in (2, 3, 4):
            ws.cell(row=r, column=cc).border = Border(bottom=Side(style="thin", color=BLUE))
        r += 1
        band = False
        continue
    if term:
        band = not band
    bg = CARD if band else "FFFFFF"
    t = ws.cell(row=r, column=2, value=term)
    t.font = font(11, True, NAVY)
    t.alignment = Alignment(vertical="top", indent=1, wrap_text=True)
    d = ws.cell(row=r, column=3, value=definition)
    d.font = font(10)
    d.alignment = Alignment(vertical="top", wrap_text=True)
    s = ws.cell(row=r, column=4, value=source)
    s.font = font(9, False, MUTED)
    s.alignment = Alignment(vertical="top", indent=1)
    for cc in (2, 3, 4):
        ws.cell(row=r, column=cc).fill = fill(bg)
    need = max(lines_for(definition, ws.px_of(3, 3) - 10), lines_for(term, ws.px_of(2, 2) - 20, 11, True),
               lines_for(source, ws.px_of(4, 4) - 20, 9))
    ws.row_dimensions[r].height = row_h(need) + 14
    r += 1
r = spacer(ws, r, 12)
c = ws.cell(row=r, column=2, value="Glossary source:")
c.font = font(9, False, MUTED)
link(ws.cell(row=r, column=3), "https://nvlpubs.nist.gov/nistpubs/ir/2013/nist.ir.7298r2.pdf",
     "NIST IR 7298 Rev. 2 – Glossary of Key Information Security Terms", size=9)
print_setup(ws)

# ================================================================== Reducing Attack Surface
ws = wb.create_sheet("Reducing Attack Surface")
setup(ws, [2.5, 5, 19, 19, 19, 19, 19, 2.5], NAVY)
banner(ws, "Reducing Attack Surface", "Fewer ways in, and harder ones")
r = 6
block(ws, r, "Attack Surface is the sum of all the points where an unauthorized user can access a system. Think of your attack "
             "surface as points of entry into your home, the threat as a thief, and the attack vector as the means by which "
             "the threat gains access to the point of entry. Our goal is to reduce the available entry points, increase the "
             "difficulty of the thief’s job, and render the thief’s attack vector useless.", size=11, pad=10)
r += 1
block(ws, r, "In this case, if we live in a neighborhood where it is likely that thieves will rob us, we may put bars on our "
             "windows. We may put extra locks or a bracing mechanism on the door. We may refuse to answer the door when "
             "someone suspicious knocks. Furthermore, we are building up our castle (HARDENING) to prevent the petty thieves "
             "from accessing our resources.", pad=10)
r += 1
block(ws, r, rich(("Maintain the balance.  ", True, NAVY),
                  ("We might make our house so secure that even WE can’t get into it. The opposite is also true, we may "
                   "allow our house to be so accessible that anyone can waltz in. How can we maintain the balance? We threat "
                   "model and end with risk mitigation!", False, INK)),
      bg=BLUE_TINT, border=ACCENT_LEFT, pad=16, indent=1, align=VCENTER)
r = spacer(ws, r + 1, 14)
r = section(ws, r, "Practical attack surface reduction")
actions = [
    ("Enable automatic updates.", "Exploiting unpatched software is now the #1 way breaches start."),
    ("Retire end-of-life devices.", "If the vendor no longer ships security updates, the device will never be patched."),
    ("Update your router.", "Install firmware updates and change the default admin password. Routers and other edge "
                            "devices are a favorite target."),
    ("Disable unnecessary services.", ""),
    ("Delete unused applications.", "Use your web browser instead of applications when possible."),
    ("Configure devices and applications correctly.", "Reference a STIG or vendor documentation."),
    ("Implement least privilege.", "Separate your accounts into one administrator and one user, and actually restrict user "
                                   "rights following separation of roles."),
    ("Use a password manager.", "Let it create a unique, strong password for every account."),
    ("Turn on MFA – or better, passkeys.", "Passkeys can’t be phished or reused. Use them wherever they’re offered."),
    ("Back up your things, frequently!", "This will protect you from mistakes AND ransomware. Keep one copy off the "
                                         "device, and test a restore now and then."),
]
for i, (head, detail) in enumerate(actions, start=1):
    b = ws.cell(row=r, column=2, value=i)
    b.font = font(11, True, "FFFFFF")
    b.fill = fill(BLUE)
    b.alignment = CENTER
    block(ws, r, rich((head, True, INK), (("  " + detail) if detail else "", False, INK2)),
          c1=3, c2=7, pad=10, indent=1, align=VCENTER)
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height, 24)
    for cc in range(3, 8):
        ws.cell(row=r, column=cc).border = ROWLINE
    r += 1
r = spacer(ws, r, 10)
block(ws, r, "These are some examples. Your device/application specific STIG will be comprehensive in helping you configure "
             "your applications and devices properly. It will likely point you in the right direction in disabling "
             "unnecessary services.", color=INK2, pad=8)
r = spacer(ws, r + 1, 14)
r = section(ws, r, "What this means for you")
block(ws, r, "You’ll notice when looking at the next tab that unpatched vulnerabilities, weak or stolen passwords, and "
             "phishing (social engineering) are the leading attack vectors. This means that if we cover those bases "
             "we’ll be much safer!", pad=8)
r += 1
c = ws.cell(row=r, column=2)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
link(c, "#'Common Attack Vectors'!A1", "Next: Common Attack Vectors  →", bold=True)
ws.row_dimensions[r].height = 22
print_setup(ws)

# ================================================================== Common Attack Vectors
ws = wb.create_sheet("Common Attack Vectors")
CAV_B = units(max(text_px(s) for s in ["Credential abuse (stolen/guessed passwords)", "Exploiting a vulnerability",
                                       "Human element involved"]) + 34)
setup(ws, [2.5, CAV_B, 11, 11, 11, 3] + [9.5] * 7 + [2.5], NAVY)
banner(ws, "Common Attack Vectors", "How attackers actually get in, and what that means for you")
r = 6
block(ws, r, "Here you’ll find some statistics regarding common attack vectors. These statistics were gathered primarily "
             "from businesses who were victims of attacks. They are still relevant to us because regardless of the "
             "percentages, among all cyber attacks there is a common theme:", pad=8)
r = spacer(ws, r + 1, 8)
themes = [("Vulnerability exploitation", "unpatched or misconfigured systems – now the #1 way in"),
          ("Phishing attacks", "social engineering, increasingly by text and phone"),
          ("Weak credentials", "reused, guessed or stolen passwords"),
          ("Ransomware", "generally delivered via social engineering attempts")]
for i, (head, sub) in enumerate(themes, start=1):
    b = ws.cell(row=r, column=2, value=rich((f"0{i}   ", True, BLUE), (head, True, NAVY)))
    b.alignment = Alignment(vertical="center", indent=1)
    b.fill = fill(CARD)
    block(ws, r, sub, c1=3, c2=13, color=INK2, bg=CARD, pad=10, align=VCENTER)
    ws.row_dimensions[r].height = 24
    for cc in range(2, 14):
        ws.cell(row=r, column=cc).border = Border(bottom=Side(style="thin", color="FFFFFF"))
    r += 1
r = spacer(ws, r, 14)
r = section(ws, r, "How we defend")
r = bullets(ws, r, [
    "We can avoid phishing, smishing, and vishing attacks by being hyper aware of email, text, and calls that are not from "
    "somebody we explicitly know. Be extra careful on your phone – people click more there.",
    "We can check to see if our accounts have been compromised on haveibeenpwned.com, ensure that we do not have weak or "
    "reused passwords (a password manager makes this easy), and enable multi-factor authentication or passkeys.",
    "We can ensure that our systems are set up to automatically update, thereby applying critical security patches before an "
    "attacker can exploit them. This will not protect you from “zero-days”, but the majority of zero-days are "
    "patched quickly and automatic updates will keep the vulnerability window as small as possible.",
    "We can ensure our devices and applications are hardened to a reasonable degree, and replace ones that no longer get "
    "security updates.",
    "We can ensure we are frequently backing up our files to prevent the effectiveness of ransomware.",
], c2=13, inline=True)
r = spacer(ws, r, 8)
block(ws, r, rich(("NOTE:  ", True, NAVY), ("The number one agent that leads to compromise is ultimately ourselves. There are "
                                             "many free cyber security awareness programs online that can help educate you. "
                                             "A little goes a long way!", False, INK)),
      c2=13, bg=BLUE_TINT, border=ACCENT_LEFT, pad=16, indent=1, align=VCENTER)
r = spacer(ws, r + 1, 14)
r = section(ws, r, "The numbers: Verizon 2026 Data Breach Investigations Report")
block(ws, r, "22,000+ confirmed data breaches in 145 countries, collected from October 2024 to November 2025.",
      c2=13, size=9, color=MUTED, italic=True, pad=4)
r = spacer(ws, r + 1, 8)
for c1, c2, value, label in [(2, 2, 0.31, "of breaches now start by exploiting a software vulnerability – the #1 way in"),
                             (3, 5, 0.48, "of breaches involve ransomware (up from 44%)"),
                             (7, 9, 0.62, "involve the human element: a mistake, a click, a scam"),
                             (10, 13, 0.40, "higher click rates on phone-based scams (texts, calls) than on email")]:
    tile(ws, r, c1, c2, value, label, fmt="0%", big=26)
ws.row_dimensions[r + 1].height = 40
r += 3

datasets = [
    ("How breaches start", "known initial access vectors, % of breaches", ["2025 DBIR", "2026 DBIR"],
     [("Exploiting a vulnerability", [20, 31]), ("Phishing", [16, 16]), ("Credential abuse (stolen/guessed passwords)", [22, 13])],
     "Source: Verizon DBIR 2025 and 2026. Pretexting (fake-scenario scams) added another 6% in 2026."),
    ("The bigger picture", "% of breaches", ["2025 DBIR", "2026 DBIR"],
     [("Ransomware present", [44, 48]), ("Human element involved", [60, 62]), ("Third party involved", [30, 48])],
     "Source: Verizon DBIR 2025 and 2026 executive summaries (2025 human element reported as “around 60%”)."),
]
for title, unit, periods, rows, source in datasets:
    top = r
    ws.cell(row=r, column=2, value=rich((title, True, NAVY), ("   " + unit, False, MUTED)))
    ws.row_dimensions[r].height = 22
    r += 1
    head = r
    for i, h in enumerate([""] + periods, start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(9, True, INK2)
        c.alignment = Alignment(horizontal="right" if i > 2 else "left", vertical="center", indent=1 if i == 2 else 0)
        c.border = Border(bottom=Side(style="medium", color=INK2))
    ws.row_dimensions[r].height = 20
    for label, values in rows:
        r += 1
        c = ws.cell(row=r, column=2, value=label)
        c.font = font(10)
        c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        c.border = ROWLINE
        for i, v in enumerate(values, start=3):
            c = ws.cell(row=r, column=i, value=v / 100)
            c.number_format = "0%"
            c.font = font(10, bold=(i == 2 + len(values)))
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.border = ROWLINE
        ws.row_dimensions[r].height = max(20, row_h(lines_for(label, ws.px_of(2, 2) - 20)) + 8)
    r += 1
    block(ws, r, source, c1=2, c2=5, size=8, color=MUTED, italic=True, pad=4, indent=1)
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.gapWidth = 60
    ch.overlap = -8
    ch.style = 2
    data = Reference(ws, min_col=3, max_col=2 + len(periods), min_row=head, max_row=head + len(rows))
    cats = Reference(ws, min_col=2, min_row=head + 1, max_row=head + len(rows))
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for sser, color in zip(ch.series, SERIES):
        sser.graphicalProperties = GraphicalProperties(solidFill=color)
        sser.graphicalProperties.line = LineProperties(solidFill="FFFFFF", w=12700)
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.crosses = "max"          # keep the % axis at the bottom when categories run top-down
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.numFmt = "0%"
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.legend.position = "b"
    ch.height, ch.width = 6.0, 14.5
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ws.add_chart(ch, f"G{top}")
    r = max(r, top + 12) + 1

r = section(ws, r, "Small and medium-sized businesses (2026 DBIR)", c2=13)
r = bullets(ws, r, [
    "Initial access: exploiting vulnerabilities 26%, credential abuse 13%, phishing 9%.",
    "Every SMB breach in the dataset came from outside attackers, and all were financially motivated.",
    "Small organizations are disproportionately hit by ransomware – the 2025 DBIR found it in 88% of SMB breaches.",
], c2=13, inline=True)
r = spacer(ws, r, 6)
c = ws.cell(row=r, column=2)
link(c, DBIR26, "Read the full 2026 DBIR  →", bold=True)
print_setup(ws, landscape=True)

# ================================================================== trackers
DEV = tracker(
    wb.create_sheet("Devices"), "Devices", "Everything you own or manage that could be a way in",
    "List every device: phones, laptops, desktops, tablets, routers, smart-home gear, NAS, consoles. Order this by security "
    "category. Pick answers from the dropdowns; hover a column header for tips.",
    [("Device", 30, None, {}, None),
     ("Security Category", 11, LMH, CAT, None),
     ("Infected", 9, YNM, {"Yes": BAD, "Maybe": WARN, "No": GOOD}, "Yes"),
     ("Backup", 9, YN, YES_GOOD, "No"),
     ("Password Protected", 10, YN, YES_GOOD, "No"),
     ("Password Complexity", 11, ["Complex", "Weak"], {"Complex": GOOD, "Weak": BAD}, "Weak"),
     ("Unique Password", 9, YN, YES_GOOD, "No"),
     ("OS Encryption", 10, YN, YES_GOOD, "No"),
     ("STIG", 8, YN, YES_GOOD, "No"),
     ("Physical Security", 9, YNM, {"Yes": GOOD, "Maybe": WARN, "No": BAD}, "No"),
     ("Automatic Updates", 10, YN, YES_GOOD, "No"),
     ("Still Supported", 10, YN, YES_GOOD, "No")],
    ["Example: personal laptop", "High", "No", "Yes", "Yes", "Complex", "Yes", "Yes", "No", "Yes", "Yes", "Yes"],
    {"Device": "Order this by security category.",
     "Security Category": "How bad would it be if this device were compromised?  Low / Moderate / High.",
     "Infected": "Any sign of malware? Unexpected pop-ups, slowness, unknown apps. “Maybe” = worth a scan.",
     "Backup": "Is the important data on this device backed up somewhere else?",
     "STIG": "Has this device been hardened to its STIG (Security Technical Implementation Guide) or the vendor’s "
             "security baseline?\n\nFind STIGs: https://www.cyber.mil/stigs/downloads/",
     "Physical Security": "Could someone walk off with it or plug into it? Locked room, cable lock, screen lock.",
     "Still Supported": "Does the vendor still ship security updates for it? End-of-life devices (old phones, routers, "
                        "Windows versions) never get patched again – plan to replace them."},
)
ACC = tracker(
    wb.create_sheet("Accounts"), "Accounts", "Every login that protects something you care about",
    "List every account: email, banking, social media, shopping, work, cloud storage, password manager. Order this by "
    "security category. Check each email address at haveibeenpwned.com.",
    [("Account", 32, None, {}, None),
     ("Security Category", 11, LMH, CAT, None),
     ("Unique Password", 10, YN, YES_GOOD, "No"),
     ("Complex Password", 10, YN, YES_GOOD, "No"),
     ("MFA", 8, YN, YES_GOOD, "No"),
     ("Account Hardening", 11, YN, YES_GOOD, "No"),
     ("Password Sharing", 11, ["Shared", "Safe"], {"Safe": GOOD, "Shared": BAD}, "Shared"),
     ("Pwned?", 13, ["Compromised", "Clean"], {"Clean": GOOD, "Compromised": BAD}, "Compromised")],
    ["Example: primary email", "High", "Yes", "Yes", "Yes", "No", "Safe", "Clean"],
    {"Account": "Order this by security category. Start with email – it can reset every other password.",
     "Security Category": "How bad would it be if this account were taken over?  Low / Moderate / High.",
     "Unique Password": "Used nowhere else. A password manager makes this painless.",
     "MFA": "Multi-factor authentication: a second step (authenticator app, security key, text code) on top of your "
            "password. A passkey counts as Yes.",
     "Account Hardening": "Recovery email/phone up to date, security questions not guessable, unused sessions signed out, "
                          "login alerts on.",
     "Pwned?": "Check your email address at https://haveibeenpwned.com. If it shows up in a breach, change that password "
               "everywhere you used it."},
)
APP = tracker(
    wb.create_sheet("Applications"), "Applications", "Now’s your time to shine!",
    "This one is for you to do! We started the columns for you: list the apps on your phone and computers, order them by "
    "security category, and delete anything you don’t need.",
    [("Application", 30, None, {}, None),
     ("Security Category", 11, LMH, CAT, None),
     ("Up to Date", 10, YN, YES_GOOD, "No"),
     ("Official Source", 10, YN, YES_GOOD, "No"),
     ("Permissions Reviewed", 11, YN, YES_GOOD, "No"),
     ("Sign-in Protected", 10, ["Yes", "No", "N/A"], {"Yes": GOOD, "No": BAD}, "No"),
     ("Still Needed", 10, YN, YES_GOOD, "No")],
    ["Example: banking app", "High", "Yes", "Yes", "Yes", "Yes", "Yes"],
    {"Application": "Order this by security category. Include browser extensions – they can see everything you do online.",
     "Official Source": "Installed from the official app store or the vendor’s own site, not a random download.",
     "Permissions Reviewed": "Location, camera, microphone, contacts – only if the app truly needs them.",
     "Sign-in Protected": "If the app has its own account: unique password plus MFA or a passkey. N/A if it has no login.",
     "Still Needed": "No? Delete it. Unused apps are attack surface with no benefit."},
)
INF = tracker(
    wb.create_sheet("Information"), "Information", "The data you can’t afford to lose or leak",
    "List the information that matters: photos, tax and financial records, IDs and passports, medical records, work files, "
    "password manager vault. Order this by security category.",
    [("Information", 32, None, {}, None),
     ("Security Category", 11, LMH, CAT, None),
     ("Where It Lives", 26, None, {}, None),
     ("Backed Up", 9, YN, YES_GOOD, "No"),
     ("Off-device Copy", 10, YN, YES_GOOD, "No"),
     ("Restore Tested", 10, YN, YES_GOOD, "No"),
     ("Encrypted", 10, YN, YES_GOOD, "No"),
     ("Access Limited", 10, YN, YES_GOOD, "No")],
    ["Example: tax & financial records", "High", "Laptop + cloud drive", "Yes", "Yes", "No", "Yes", "Yes"],
    {"Information": "Order this by security category.",
     "Where It Lives": "Device, cloud service or physical location (type it in).",
     "Off-device Copy": "At least one backup that isn’t on the same device – an external drive kept elsewhere or "
                        "a cloud backup. Ransomware and dead drives take out local copies.",
     "Restore Tested": "Have you actually restored a file from the backup? An untested backup is a hope, not a plan.",
     "Encrypted": "Stored on an encrypted device or in an encrypted service/vault.",
     "Access Limited": "Only the people who need it can open it. Shared links and old family accounts count."},
)
TRACKERS = [DEV, ACC, APP, INF]

# ================================================================== Resources
ws = wb.create_sheet("Resources")
groups = [
    ("Threat reports (the stats on Common Attack Vectors)", [
        ("Verizon 2026 Data Breach Investigations Report (DBIR)", DBIR26),
        ("Verizon 2026 DBIR executive summary",
         "https://www.verizon.com/business/resources/executivebriefs/2026-dbir-executive-summary.pdf"),
        ("Verizon DBIR home (all editions)", "https://www.verizon.com/business/resources/reports/dbir/"),
    ]),
    ("Protect yourself", [
        ("CISA Secure Our World – simple steps for everyone", "https://www.cisa.gov/secure-our-world"),
        ("CISA – Use strong passwords (and a password manager)", "https://www.cisa.gov/secure-our-world/use-strong-passwords"),
        ("FIDO Alliance – What are passkeys?", "https://fidoalliance.org/passkeys/"),
        ("Have I Been Pwned – check if your email was in a breach", "https://haveibeenpwned.com/"),
    ]),
    ("Standards & frameworks", [
        ("NIST IR 7298 Rev. 2 – Glossary of Key Information Security Terms",
         "https://nvlpubs.nist.gov/nistpubs/ir/2013/nist.ir.7298r2.pdf"),
        ("NIST SP 800-53 Rev. 5 – Security and Privacy Controls", "https://csrc.nist.gov/pubs/sp/800/53/r5/upd1/final"),
        ("ISO 31000 – Risk management guidelines", "https://www.iso.org/obp/ui/#iso:std:iso:31000:ed-2:v1:en"),
    ]),
    ("Hardening guides (STIGs)", [
        ("DISA STIG document library", "https://www.cyber.mil/stigs/downloads/"),
        ("DISA SRG & STIG library compilations", "https://www.cyber.mil/stigs/compilations/"),
        ("STIG Viewer (browse STIGs online)", "https://www.stigviewer.com/stigs"),
    ]),
    ("Further reading", [
        ("Ars Technica – How I learned to stop worrying (mostly) and love my threat model",
         "https://arstechnica.com/information-technology/2017/07/how-i-learned-to-stop-worrying-mostly-and-love-my-threat-model/"),
    ]),
    ("Earlier editions (the 2022 version of this workbook)", [
        ("Verizon 2022 DBIR",
         "https://www.verizon.com/business/resources/reports/2022/dbir/2022-data-breach-investigations-report-dbir.pdf"),
        ("Kroll Q1 2022 Threat Landscape",
         "https://www.kroll.com/en/insights/publications/cyber/q1-2022-threat-landscape-threat-actors-target-email-access-extortion"),
        ("Kaspersky Incident Response Analyst Report 2021",
         "https://media.kasperskycontenthub.com/wp-content/uploads/sites/43/2021/09/13085018/Incident-Response-Analyst-Report-eng-2021.pdf"),
        ("Kaspersky – The nature of cyber incidents (2022)",
         "https://media.kasperskycontenthub.com/wp-content/uploads/sites/43/2022/09/02120838/Kaspersky-The-nature-of-cyber-incidents_v11-1.pdf"),
    ]),
]
RES_B = min(66, units(max(text_px(n, 10, True) for _, items in groups for n, _ in items) + 30))
setup(ws, [2.5, RES_B, 46, 16, 2.5], "898781")
banner(ws, "Resources", "Where the numbers and ideas in this workbook come from")
r = 6
for gname, items in groups:
    r = section(ws, r, gname, c2=4)
    for name, url in items:
        n = ws.cell(row=r, column=2, value=name)
        n.font = font(10, True, INK)
        n.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        u = ws.cell(row=r, column=3)
        short = url.split("//", 1)[1]
        short = short if len(short) <= 70 else short[:67] + "…"
        link(u, url, short, size=9)
        u.alignment = VCENTER
        for cc in (2, 3, 4):
            ws.cell(row=r, column=cc).border = ROWLINE
        ws.row_dimensions[r].height = max(24, row_h(lines_for(name, ws.px_of(2, 2) - 20, 10, True)) + 10)
        r += 1
    r = spacer(ws, r, 12)
print_setup(ws)

# ================================================================== Start Here (filled last)
ws = start
setup(ws, [2.5, 17, 17, 17, 17, 17, 17, 2.5], NAVY)
banner(ws, "Threat Modeling + Risk Management Framework",
       "Asset & Mitigation Checklist  ·  FNIT Guy, your Friendly Neighborhood IT Guy  ·  2026 edition", back=False)
r = 6
block(ws, r, "In this sheet, you’ll find a succinct reference to assist you in protecting what’s important to you. "
             "This isn’t meant to be comprehensive but to act as a launching pad and foundation for your own research. "
             "You’ll see in some of the tabs there is a column labeled “STIG” AKA Security Technical Implementation "
             "Guide. This is where you’ll find very specific hardening guidance for your specific application or device. "
             "You’ll see a link listed in the comments for those columns, use them!",
      bg=BLUE_TINT, border=ACCENT_LEFT, pad=16, indent=1, align=VCENTER)
r = spacer(ws, r + 1, 14)

r = section(ws, r, "How to use this workbook")
steps = [("1", "Learn the basics", "Skim Definitions, Reducing Attack Surface and Common Attack Vectors."),
         ("2", "Take inventory", "Fill in Devices, Accounts, Applications and Information. Pick answers from the dropdowns."),
         ("3", "Fix the red", "Every tracker has a Priority column. Knock out “Fix first” items, then work down.")]
for i, (n, head, body) in enumerate(steps):
    c1 = 2 + i * 2
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + 1)
    ws.merge_cells(start_row=r + 1, start_column=c1, end_row=r + 1, end_column=c1 + 1)
    h = ws.cell(row=r, column=c1, value=rich((n + "  ", True, BLUE), (head, True, NAVY), size=11))
    h.alignment = Alignment(vertical="center", indent=2)
    b = ws.cell(row=r + 1, column=c1, value=body)
    b.font = font(10, False, INK2)
    b.alignment = Alignment(wrap_text=True, vertical="top", indent=2)
    for rr in (r, r + 1):
        for cc in (c1, c1 + 1):
            ws.cell(row=rr, column=cc).fill = fill(CARD)
ws.row_dimensions[r].height = 26
ws.row_dimensions[r + 1].height = row_h(max(lines_for(b, ws.px_of(2, 3) - 20) for _, _, b in steps)) + 14
r += 2
r = spacer(ws, r, 6)
block(ws, r, rich(("Legend:  ", True, INK2), ("white cells are yours to fill in; gray italic rows are examples.  ", False, INK2),
                  ("Green", True, GOOD[1]), (" = good  ·  ", False, INK2), ("Yellow", True, WARN[1]),
                  (" = check it  ·  ", False, INK2), ("Red", True, BAD[1]), (" = gap to fix", False, INK2)), pad=6)
r = spacer(ws, r + 1, 14)

r = section(ws, r, "Your scorecard")
rng = lambda t, L: f"'{t['name']}'!{L}{FIRST}:{L}{LAST}"
tracked = "+".join(f"COUNTA({rng(t, 'B')})" for t in TRACKERS)
gaps = "+".join(f"SUM({rng(t, t['gaps'])})" for t in TRACKERS)
fix_first = "+".join(f'COUNTIF({rng(t, t["priority"])},"Fix first")' for t in TRACKERS)
done = "+".join(f'COUNTIF({rng(t, t["priority"])},"Done")' for t in TRACKERS)
pct = lambda t, L: (f'=IF(COUNTA({rng(t, "B")})=0,"–",COUNTIF({rng(t, L)},"Yes")/COUNTA({rng(t, "B")}))')
tiles = [
    (f"={tracked}", "0", "Items tracked across all four checklists"),
    (f"={gaps}", "0", "Open gaps"),
    (f"={fix_first}", "0", "“Fix first” items (High category with gaps)"),
    (f'=IF(({tracked})=0,"–",({done})/({tracked}))', "0%", "Items fully done (no gaps)"),
    (pct(ACC, "F"), "0%", "Accounts with MFA or passkeys"),
    (pct(INF, "F"), "0%", "Important info with an off-device backup"),
]
for k, (formula, fmt, label) in enumerate(tiles):
    row = r + (k // 3) * 3
    c1 = 2 + (k % 3) * 2
    tile(ws, row, c1, c1 + 1, formula, label, fmt)
    ws.row_dimensions[row + 2].height = 8
r += 6
block(ws, r, "Updates automatically as you fill in the checklists (example rows aren’t counted).",
      size=9, color=MUTED, italic=True, pad=4)
r = spacer(ws, r + 1, 14)

r = section(ws, r, "First questions")
for q in ["What do I want to protect?", "Who do I want to protect it from?",
          "How likely is it that I’ll need to protect it?", "How bad are the consequences if I fail?",
          "How much trouble am I willing to go through to prevent those consequences?"]:
    block(ws, r, q, size=11, bg=BLUE_TINT, border=ACCENT_LEFT, pad=10, indent=1, align=VCENTER)
    r += 1
r = spacer(ws, r, 14)

r = section(ws, r, "The framework")
framework = [
    ("1", "Asset inventory", "What do we need to protect?", []),
    ("2", "Security category", "Accounts  ·  Devices  ·  Applications  ·  Information", []),
    ("3", "Threat inventory", "Threat source?  Identify threats  ·  Who/What  ·  Attack vectors", []),
    ("4", "Assess vulnerabilities", None, [
        ("Accounts", "reduce attack surface, make it very difficult to gain unauthorized access",
         "Poor password security  ·  Password reuse  ·  Weak passwords  ·  No MFA"),
        ("Devices", "reduce attack surface, harden what’s left",
         "Poor password security  ·  Weak passwords  ·  No MFA  ·  Passwords stored in clear text  ·  "
         "Weak baseline security configurations  ·  No hardening  ·  Unnecessary services enabled  ·  "
         "Unnecessary applications downloaded  ·  Lack of network security  ·  No longer supported"),
        ("Applications", None, "Weak configurations  ·  No security  ·  Out of date"),
        ("Information", None, "No backup  ·  No encryption")]),
    ("5", "Identify and implement controls to mitigate risk",
     "Threat + vulnerability = risk.  Are there threats?  Are there vulnerabilities?  Is there risk? Yes.  Controls to fix.", []),
]
for n, head, detail, subs in framework:
    badge = ws.cell(row=r, column=2, value=n)
    badge.font = font(14, True, "FFFFFF")
    badge.fill = fill(NAVY)
    badge.alignment = CENTER
    t = ws.cell(row=r, column=3, value=head)
    t.font = font(12, True, NAVY)
    t.alignment = Alignment(vertical="center", indent=2)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 26
    r += 1
    if detail:
        block(ws, r, detail, c1=3, c2=7, color=INK2, pad=6, indent=1)
        r += 1
    for name, why, items in subs:
        block(ws, r, rich((name, True, INK), ((" – " + why) if why else "", False, INK2)), c1=3, c2=7, pad=2, indent=1)
        r += 1
        block(ws, r, items, c1=3, c2=7, color=MUTED, pad=6, indent=3)
        r += 1
    r = spacer(ws, r, 8)
r = spacer(ws, r, 8)

r = section(ws, r, "What’s inside")
directory = [
    ("Definitions", "Threats, vulnerabilities, phishing, passkeys and more, in plain English."),
    ("Reducing Attack Surface", "The house-and-thief analogy, plus ten practical ways to shrink your attack surface."),
    ("Common Attack Vectors", "How attackers actually get in, with the 2026 numbers behind it."),
    ("Devices", "Checklist: every device you own or manage."),
    ("Accounts", "Checklist: every online account."),
    ("Applications", "Checklist: the apps and extensions you use."),
    ("Information", "Checklist: the data you can’t afford to lose or leak."),
    ("Resources", "Reports, standards, hardening guides and further reading."),
]
for name, desc in directory:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = ws.cell(row=r, column=2)
    link(c, f"#'{name}'!A1", name + "  →", bold=True)
    c.alignment = Alignment(vertical="center", indent=1)
    block(ws, r, desc, c1=4, c2=7, color=INK2, pad=10, align=VCENTER)
    for cc in range(2, 8):
        ws.cell(row=r, column=cc).border = ROWLINE
    r += 1
r = spacer(ws, r, 18)
c = ws.cell(row=r, column=2)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
link(c, VIDEO, "Watch the video: Threat Modeling for Normal People", bold=True)
c2 = ws.cell(row=r, column=5)
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
link(c2, SITE, "fnitguy.tech")
c2.alignment = Alignment(horizontal="right")
print_setup(ws)

# Indents only render in LibreOffice when horizontal alignment is explicitly "left".
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            a = cell.alignment
            if a is not None and a.indent and a.horizontal in (None, "general"):
                cell.alignment = Alignment(horizontal="left", vertical=a.vertical, wrap_text=a.wrap_text, indent=a.indent)

wb.calculation.fullCalcOnLoad = True   # formulas are written without cached values; compute on open
wb.active = 0
wb.save(OUT)
print("saved", OUT)
